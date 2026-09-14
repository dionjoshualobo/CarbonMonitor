"""Parse uploaded emission CSVs/XLSX files."""

import csv
import io
from datetime import datetime, timedelta, timezone
from typing import Any
from uuid import UUID

from sqlalchemy.orm import Session

REQUIRED_COLUMNS = {"timestamp", "metric", "value"}
ALLOWED_METRICS = {"electricity", "diesel", "petrol", "lpg"}


def _parse_timestamp(value: str | None) -> str:
    if not value or not value.strip():
        raise ValueError("missing timestamp")
    try:
        return datetime.fromisoformat(value.strip().replace("Z", "+00:00")).isoformat()
    except ValueError as exc:
        raise ValueError(f"invalid timestamp '{value}'") from exc


def parse_emissions_csv(content: bytes) -> list[dict[str, Any]]:
    reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig")))
    headers = {header.strip() for header in (reader.fieldnames or [])}
    missing = REQUIRED_COLUMNS - headers
    if missing:
        raise ValueError(f"CSV missing required columns: {', '.join(sorted(missing))}")

    rows: list[dict[str, Any]] = []
    for line_no, raw in enumerate(reader, start=2):
        try:
            value = float(raw["value"])
        except (TypeError, ValueError) as exc:
            raise ValueError(f"row {line_no}: 'value' must be numeric") from exc
        metric = (raw.get("metric") or "").strip()
        if not metric:
            raise ValueError(f"row {line_no}: 'metric' is required")
        if metric not in ALLOWED_METRICS:
            raise ValueError(f"row {line_no}: unknown metric '{metric}' — expected one of {', '.join(sorted(ALLOWED_METRICS))}")
        try:
            timestamp = _parse_timestamp(raw.get("timestamp"))
        except ValueError as exc:
            raise ValueError(f"row {line_no}: {exc}") from exc
        rows.append(
            {
                "timestamp": timestamp,
                "metric": metric,
                "value": value,
                "unit": (raw.get("unit") or "").strip() or None,
                "facility_name": (raw.get("facility_name") or "").strip() or None,
                "source": "upload",
            }
        )
    if not rows:
        raise ValueError("CSV contains no data rows")
    return rows


def convert_xlsx_to_csv_bytes(content: bytes) -> bytes:
    """Convert xlsx (MescomBill format) to CSV with timestamp/metric/value/unit/facility_name."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["timestamp", "metric", "value", "unit", "facility_name"])

    epoch = datetime(1899, 12, 30, tzinfo=timezone.utc)
    for row in ws.iter_rows(min_row=2, values_only=True):
        month_val = row[0]  # column A: datetime or serial number
        total_units = row[19]  # column T: Total Units (0-indexed)
        if month_val is None or total_units is None:
            continue
        # openpyxl may return datetime or int/float serial
        if hasattr(month_val, "isoformat"):
            dt = month_val.replace(tzinfo=timezone.utc) if month_val.tzinfo is None else month_val
        else:
            dt = epoch + timedelta(days=int(month_val))
        writer.writerow([dt.isoformat(), "electricity", total_units, "kWh", ""])

    return buf.getvalue().encode("utf-8")


def index_upload_readings(db: Session, upload_id: str, facility_id: UUID | None, rows: list[dict[str, Any]]) -> int:
    """Persist parsed rows as ActivityRecords + CalculatedEmissions."""
    from models.activity_record import ActivityRecord
    from models.emission_factor import EmissionFactor
    from services.calculation import calculate_emissions

    if facility_id is None:
        raise ValueError("facility_id is required — user has no facility assigned")

    def next_month_start(dt: datetime) -> datetime:
        if dt.month == 12:
            return dt.replace(year=dt.year + 1, month=1, day=1)
        return dt.replace(month=dt.month + 1, day=1)

    # cache factors per activity_type
    factor_cache: dict[str, Any] = {}

    def get_factor(activity_type: str):
        if activity_type not in factor_cache:
            f = (
                db.query(EmissionFactor)
                .filter(
                    EmissionFactor.activity_type == activity_type,
                    EmissionFactor.region.is_(None),
                )
                .first()
            )
            if f is None:
                raise ValueError(f"no emission_factor for activity_type={activity_type}")
            factor_cache[activity_type] = f
        return factor_cache[activity_type]

    count = 0
    for row in rows:
        period_start = datetime.fromisoformat(row["timestamp"])
        metric = row["metric"]
        factor = get_factor(metric)
        ar = ActivityRecord(
            facility_id=facility_id,
            period_start=period_start,
            period_end=next_month_start(period_start),
            activity_type=metric,
            quantity=row["value"],
            unit=row.get("unit") or "kWh",
            source="csv",
            confirmed_by_user=False,
        )
        db.add(ar)
        db.flush()
        ce = calculate_emissions(ar, factor)
        db.add(ce)
        count += 1

    db.commit()
    return count
